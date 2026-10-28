import os
import tempfile
import pytest
from docx import Document
from app.parser.docx_extractor import parse_docx_file
from app.parser.language_detector import detect_language
from app.parser.code_sanitizer import partition_code_and_output, strip_line_numbering


def test_language_detection():
    cpp_code = """
    #include <iostream>
    using namespace std;
    int main() {
        cout << "Hello DSA" << endl;
        return 0;
    }
    """
    lang, conf = detect_language(cpp_code)
    assert lang == "cpp"
    assert conf > 70

    py_code = """
    def bubble_sort(arr):
        n = len(arr)
        for i in range(n):
            for j in range(0, n-i-1):
                if arr[j] > arr[j+1]:
                    arr[j], arr[j+1] = arr[j+1], arr[j]
        return arr
    print(bubble_sort([5, 2, 8, 1]))
    """
    lang, conf = detect_language(py_code)
    assert lang == "python"
    assert conf > 70

    sql_code = """
    CREATE TABLE Students (
        id INT PRIMARY KEY,
        name VARCHAR(100),
        cgpa FLOAT
    );
    SELECT * FROM Students WHERE cgpa > 8.5;
    """
    lang, conf = detect_language(sql_code)
    assert lang == "sql"
    assert conf > 70

    r_code = """
    x <- c(10, 20, 30, 40)
    y <- c(15, 25, 35, 45)
    plot(x, y, main="Scatterplot", col="blue", pch=19)
    df <- data.frame(name=c("A", "B"), score=x[1:2])
    summary(df)
    """
    lang, conf = detect_language(r_code)
    assert lang == "r"
    assert conf > 70


def test_strip_line_numbering():
    numbered_lines = [
        "1. #include <stdio.h>",
        "2. int main() {",
        "3.     printf(\"Test\");",
        "4.     return 0;",
        "5. }"
    ]
    cleaned = strip_line_numbering(numbered_lines)
    assert cleaned[0] == "#include <stdio.h>"
    assert cleaned[1] == "int main() {"
    assert cleaned[2] == "    printf(\"Test\");"


def test_partition_code_and_output():
    mixed = [
        "#include <iostream>",
        "int main() { return 0; }",
        "Output:",
        "Compilation successful",
        "Array elements: 10 20 30"
    ]
    code, output = partition_code_and_output(mixed)
    assert len(code) == 2
    assert "#include <iostream>" in code
    assert len(output) == 3
    assert "Array elements: 10 20 30" in output


def test_interleaved_multi_cycle_code_output():
    """Verify that multiple code-output-code cycles are preserved (e.g. EXP #2 line graph + scatterplot)."""
    raw_lines = [
        "# Line Graph",
        "x <- c(1, 2, 3, 4, 5)",
        "plot(x, type='o', col='blue')",
        "Output:",
        "[Graph window rendered]",
        "# Scatterplot",
        "y <- c(10, 20, 30, 40, 50)",
        "plot(x, y, pch=19)",
        "Output:",
        "[Scatterplot rendered]",
        "# Boxplot",
        "boxplot(x, col='green')"
    ]
    code, output = partition_code_and_output(raw_lines)
    code_text = "\n".join(code)
    output_text = "\n".join(output)

    # All 3 sub-experiments must be in code
    assert "Line Graph" in code_text
    assert "plot(x, type='o', col='blue')" in code_text
    assert "Scatterplot" in code_text
    assert "plot(x, y, pch=19)" in code_text
    assert "Boxplot" in code_text
    assert "boxplot(x, col='green')" in code_text

    # Outputs must be segregated
    assert "[Graph window rendered]" in output_text
    assert "[Scatterplot rendered]" in output_text
    assert "[Graph window rendered]" not in code_text
    assert "[Scatterplot rendered]" not in code_text


def test_repl_prompt_sanitization():
    """Verify R console prompts (>, +) and console outputs ([1]) are handled properly."""
    raw_lines = [
        "> age <- c(22, 25, 28, 32)",
        "> weight <- c(65, 70, 72, 80)",
        "> cor(age, weight)",
        "[1] 0.99868",
        "> summary_stats <- function(v) {",
        "+   list(mean = mean(v), sd = sd(v))",
        "+ }"
    ]
    code, output = partition_code_and_output(raw_lines)
    code_text = "\n".join(code)
    output_text = "\n".join(output)

    # Prompts must be stripped
    assert "age <- c(22, 25, 28, 32)" in code_text
    assert "weight <- c(65, 70, 72, 80)" in code_text
    assert "cor(age, weight)" in code_text
    assert "summary_stats <- function(v) {" in code_text
    assert "list(mean = mean(v), sd = sd(v))" in code_text
    assert not any(l.startswith("> ") for l in code)
    assert not any(l.startswith("+ ") for l in code)

    # Vector output index must be in output
    assert "[1] 0.99868" in output_text
    assert "[1] 0.99868" not in code_text


def test_full_docx_multi_practical_parsing():
    # Programmatically create a sample multi-experiment docx
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = Document()
        doc.add_heading("Computer Science Lab Manual", level=0)

        # Practical 1: C++ Linked List
        doc.add_heading("Practical 1: Singly Linked List Implementation", level=1)
        doc.add_paragraph("Aim: To implement singly linked list insertion and deletion in C++.")
        doc.add_paragraph("Theory: A linked list is a linear data structure of nodes.")
        doc.add_paragraph("#include <iostream>")
        doc.add_paragraph("using namespace std;")
        doc.add_paragraph("struct Node { int data; Node* next; };")
        doc.add_paragraph("int main() {")
        doc.add_paragraph("    cout << \"Linked List Created\" << endl;")
        doc.add_paragraph("    return 0;")
        doc.add_paragraph("}")
        doc.add_paragraph("Output:")
        doc.add_paragraph("Linked List Created")
        doc.add_paragraph("Process finished with exit code 0")

        # Practical 2: Python Binary Search
        doc.add_heading("Practical 2: Binary Search in Python", level=1)
        doc.add_paragraph("Aim: Write a python program to perform binary search on a sorted list.")
        doc.add_paragraph("def binary_search(arr, target):")
        doc.add_paragraph("    low, high = 0, len(arr) - 1")
        doc.add_paragraph("    while low <= high:")
        doc.add_paragraph("        mid = (low + high) // 2")
        doc.add_paragraph("        if arr[mid] == target: return mid")
        doc.add_paragraph("        elif arr[mid] < target: low = mid + 1")
        doc.add_paragraph("        else: high = mid - 1")
        doc.add_paragraph("    return -1")
        doc.add_paragraph("print(binary_search([1, 3, 5, 7, 9], 5))")
        doc.add_paragraph("Sample Output:")
        doc.add_paragraph("Element found at index 2")

        doc.save(tmp_path)

        # Parse with engine
        res = parse_docx_file(tmp_path)

        assert res["total_practicals_found"] == 2
        p1 = res["practicals"][0]
        p2 = res["practicals"][1]

        # Assertions for Practical 1
        assert p1["practical_number"] == 1
        assert "Singly Linked List" in p1["title"] or "Aim" in p1["title"]
        assert "#include <iostream>" in p1["code"]
        assert "cout << \"Linked List Created\"" in p1["code"]
        # CRITICAL: Verify Output is NOT in the code!
        assert "Process finished with exit code 0" not in p1["code"]
        assert "Linked List Created" in p1["discarded_output"]
        assert p1["language"] == "cpp"

        # Assertions for Practical 2
        assert p2["practical_number"] == 2
        assert "Binary Search" in p2["title"]
        assert "def binary_search" in p2["code"]
        # CRITICAL: Verify Sample Output is NOT in the code!
        assert "Element found at index 2" not in p2["code"]
        assert "Element found at index 2" in p2["discarded_output"]
        assert p2["language"] == "python"

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_quote_normalization():
    from app.parser.code_sanitizer import normalize_typographic_quotes
    code_with_smart_quotes = 'print(“Hello, World!”)\nname = ‘Alice’'
    cleaned = normalize_typographic_quotes(code_with_smart_quotes)
    assert cleaned == 'print("Hello, World!")\nname = \'Alice\''
    assert "“" not in cleaned
    assert "”" not in cleaned
    assert "‘" not in cleaned
    assert "’" not in cleaned


def test_table_merged_cell_deduplication():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = Document()
        doc.add_heading("Experiment 1: Table Test", level=1)
        doc.add_paragraph("Aim: Test table de-duplication")
        table = doc.add_table(rows=2, cols=2)
        # Merge row 0 cell 0 with cell 1
        a = table.cell(0, 0)
        b = table.cell(0, 1)
        a.merge(b)
        a.text = "Merged Header Text"
        table.cell(1, 0).text = "Row 2 Col 1"
        table.cell(1, 1).text = "Row 2 Col 2"
        doc.save(tmp_path)

        from app.parser.docx_extractor import extract_raw_lines_from_docx
        lines = extract_raw_lines_from_docx(tmp_path)
        all_text = "\n".join(lines)
        # Merged header must occur exactly once
        assert all_text.count("Merged Header Text") == 1
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

